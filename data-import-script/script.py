from openpyxl import load_workbook
import odoorpc

# Prepare the connection to the server
odoo = odoorpc.ODOO('localhost', port=8069)

# Login
odoo.login('T-snitch', 'admin', '123')


def import_work_center():
    wb = load_workbook(filename='/home/jayesh/Downloads/Work Centers.xlsx')
    sheet = wb.active
    for record in sheet.iter_rows(min_row=2, max_row=None, min_col=None, max_col=None, values_only=True):
        tags = record[1].split(',')
        tag_list = []
        for tag in tags:
            tag_id = odoo.env['mrp.workcenter.tag'].search([('name', '=', tag)])
            if tag_id:
                tag_list.append(tag_id[0])
            else:
                tag_id = odoo.env['mrp.workcenter.tag'].create({
                    'name': tag
                })
                tag_list.append(tag_id[0])

        search = odoo.env['mrp.workcenter'].search([('name', '=', record[0])])
        if not search:
            odoo.env['mrp.workcenter'].create({
                'name': record[0],
                'tag_ids': [(6, 0, tag_list)],
                'code': record[2],
            })

    for record in sheet.iter_rows(min_row=2, max_row=None, min_col=None, max_col=None, values_only=True):
        alr_workcenter_list = []
        if record[3]:
            alr_workcenters = record[3].split(',')
            for center in alr_workcenters:
                center_id = odoo.env['mrp.workcenter'].search([('name', '=', center)])
                alr_workcenter_list.append(center_id[0])
                print(alr_workcenter_list)

        if len(alr_workcenter_list) > 0:
            wc = odoo.env['mrp.workcenter'].search([('name', '=', record[0])])
            center = odoo.env['mrp.workcenter'].browse(wc)
            center.write({
                'alternative_workcenter_ids': [(6, 0, alr_workcenter_list)],
            })


import_work_center()

