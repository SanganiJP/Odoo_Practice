from openpyxl import load_workbook
import odoorpc

# Prepare the connection to the server
odoo = odoorpc.ODOO('localhost', port=8069)

# Login
odoo.login('HMS', 'admin', 'admin')


def import_product_categories():
    wb = load_workbook(filename='/home/jayesh/Downloads/Product Categories.xlsx')
    sheet = wb.active
    for record in sheet.iter_rows(min_row=2, max_row=None, min_col=None, max_col=None, values_only=True):
        parent_lst = []
        if record[1]:
            parent_cat = record[1].split(' / ')
            if len(parent_cat) == 1:
                parent_category = odoo.env['product.category'].search([('name', '=', record[1])])
                if parent_category:
                    parent_lst.append(parent_category[0])
                else:
                    id = odoo.env['product.category'].create({
                        'name': record[0],
                    })
                    parent_lst.append(id)

            if len(parent_cat) > 1:
                parent_category = odoo.env['product.category'].search([('name', '=', parent_cat[len(parent_cat) - 1])])
                sub_parent_id = odoo.env['product.category'].search([('name', '=', parent_cat[len(parent_cat) - 2])])
                if parent_category:
                    parent_lst.append(parent_category[0])
                else:
                    id = odoo.env['product.category'].create({
                        'name': parent_cat[len(parent_cat) - 1],
                        'parent_id': sub_parent_id[0]
                    })
                    parent_lst.append(id)

        if record[0]:
            cat_list = record[0].split(' / ')
            cat_lst_length = len(cat_list)
            if cat_lst_length == 1:
                category = odoo.env['product.category'].search([('name', '=', record[0])])
                if category:
                    parent_lst.append(category[0])
                else:
                    id = odoo.env['product.category'].create({
                        'name': record[0],
                    })
                    parent_lst.append(id[0])

            if cat_lst_length > 1:
                category = odoo.env['product.category'].search([('name', '=', cat_list[cat_lst_length - 1])])
                if not category:
                    odoo.env['product.category'].create({
                        'name': cat_list[cat_lst_length - 1],
                        'parent_id': parent_lst[0]
                    })


import_product_categories()
