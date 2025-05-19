import base64
from io import StringIO

import openpyxl
from openpyxl import load_workbook
import tempfile
import base64

from odoo import models, fields, api
from odoo.exceptions import ValidationError, UserError


class UpdateMrpComponentWizard(models.TransientModel):
    _name = "update.mrp.component.wizard"
    _description = "Update mrp component wizard"

    operation_type = fields.Selection(
        [('update_component', 'Update Component'), ('replace_component', 'Replace Component')], string="Operation Type",
        default="update_component")
    excel_file_for_import = fields.Binary(string='File for upload')
    product_ids = fields.One2many("mo.component.data", "mo_product_id", string="Components")
    serial_product_ids = fields.One2many("mo.component.serial.product", "mo_product_id", string="Components")
    is_active = fields.Boolean()

    def action_update_mo_component(self):
        records = self.env['mo.component.serial.product'].search([('mo_product_id', '=', self.id)])

        for rec in records:
            product_id = self.env['product.product'].search([('default_code', '=', rec.product_code)])
            old_serial_id = self.env['stock.lot'].search([('name', '=', rec.old_serial), ('product_id', '=', product_id.id)])
            new_serial_id = self.env['stock.lot'].search([('name', '=', rec.new_serial), ('product_id', '=', product_id.id)])
            mo_reord = self.env['mrp.production'].search([('name', '=', rec.mo_number)])
            if mo_reord:
                for line in mo_reord.move_raw_ids:
                    line.lot_ids = (3, old_serial_id.id)
                    line.lot_ids = (4, new_serial_id.id)


    def action_replace_mo_component(self):
        records = self.env['mo.component.data'].search([('mo_product_id', '=', self.id)])
        mo_rec_id = 0
        for rec in records:
            mo_reord = self.env['mrp.production'].search([('name', '=', rec.mo_number)])
            if mo_rec_id != mo_reord.id:
                if rec.state != 'draft':
                    continue
                else:
                    data = []
                    same_mo_records = self.env['mo.component.data'].search(
                        [('mo_product_id', '=', self.id), ('mo_number', '=', rec.mo_number)])
                    for rec in same_mo_records:
                        data.append((0, 0, {
                            'product_id': rec.new_product.id,
                            # 'product_uom_qyt': 1
                        }))
                    mo_reord.move_raw_ids = False
                    mo_reord.move_raw_ids = data
                    mo_rec_id = mo_reord.id

    def action_read(self):
        if self.excel_file_for_import:
            self.product_ids = False
            self.serial_product_ids = False
            self.is_active = True
            file_data = base64.b64decode(self.excel_file_for_import)
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp:
                temp.write(file_data)
                temp.seek(0)
                # Load workbook
                workbook = openpyxl.load_workbook(temp.name)
                sheet = workbook.active
                lst = []
                # Iterate over rows
                for row in sheet.iter_rows(min_row=2, values_only=True):  # skip header
                    if self.operation_type == 'replace_component':
                        current_product_id = self.env['product.product'].search([('default_code', '=', row[1])])
                        new_product_id = self.env['product.product'].search([('default_code', '=', row[2])])
                        lst.append((0, 0, {
                            'mo_number': row[0],
                            'current_product': current_product_id.id,
                            'new_product': new_product_id.id,
                            'state': row[3],
                            'mo_product_id': self.id
                        }))
                    if self.operation_type == 'update_component':
                        lst.append((0, 0, {
                            'mo_number': row[0],
                            'product_code': row[1],
                            'old_serial': row[2],
                            'new_serial': row[3],
                        }))

                if self.operation_type == 'replace_component':
                    self.product_ids = lst

                if self.operation_type == 'update_component':
                    self.serial_product_ids = lst

                return {
                    'type': 'ir.actions.act_window',
                    'res_model': 'update.mrp.component.wizard',
                    'view_mode': 'form',
                    'res_id': self.id,
                    'target': 'new',
                }

        else:
            raise UserError('Please select file to upload!')